#!/usr/bin/env python3
"""Build the focused Korean first-hour factorial paper as DOCX and PDF."""
from __future__ import annotations

import hashlib
import json
import re
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/FIRST_HOUR_FACTORIAL_PAPER_20261030.md"
EVIDENCE = ROOT / "docs/evidence/first_hour_factorial_20261030"
OUTPUT = EVIDENCE / "paper"
NAME = "FIRST_HOUR_FACTORIAL_PAPER_20261030"
FIGURES = tuple((EVIDENCE / "figures" / name) for name in (
    "hour_12_13_paired_precipitation.png",
    "hour_12_13_skill_factorial.png",
    "factorial_design_and_metric.png",
))
SCORE_INPUTS = tuple((EVIDENCE / "analysis" / name) for name in (
    "categorical_scores.csv", "continuous_scores.csv", "score_contrasts.csv",
    "pointwise_precipitation_contrasts.csv", "hour_13_arrays.npz", "receipt.json",
))
FIGURE_MANIFEST = EVIDENCE / "figures" / "figure_manifest.json"


def run(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(args, check=True, text=True, capture_output=True)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def set_style_font(style) -> None:
    style.font.name = "NanumMyeongjo"
    style.font.size = Pt(10)
    properties = style.element.get_or_add_rPr()
    fonts = properties.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.insert(0, fonts)
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), "NanumMyeongjo")


def format_docx(path: Path) -> None:
    doc = Document(path)
    for style in doc.styles:
        if style.type == 1:
            set_style_font(style)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2.0)
        section.left_margin = section.right_margin = Cm(2.1)
    for paragraph in doc.paragraphs:
        if re.match(r"^표 \d+\.", paragraph.text):
            paragraph.paragraph_format.keep_with_next = True
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.keep_with_next = True
    for table in doc.tables:
        for row_index, row in enumerate(table.rows):
            row_properties = row._tr.get_or_add_trPr()
            if row_properties.find(qn("w:cantSplit")) is None:
                row_properties.append(OxmlElement("w:cantSplit"))
            if row_index < len(table.rows) - 1:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
    title = doc.paragraphs[0]
    title.alignment = 1
    for run in title.runs:
        run.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = None
    doc.save(path)


def internal_inputs(markdown: str) -> list[Path]:
    links = re.findall(r"!?\[[^]]*\]\(([^)]+)\)", markdown)
    paths = []
    for link in links:
        if "://" in link or link.startswith("#"):
            continue
        path = (SOURCE.parent / link.split("#", 1)[0]).resolve()
        if not path.is_file():
            raise FileNotFoundError(f"paper link does not resolve: {link}")
        if link.endswith((".png", ".jpg", ".jpeg")):
            paths.append(path)
    return paths


def check_docx(path: Path) -> tuple[int, int]:
    with ZipFile(path) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
        ns = {
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }
        math_count = len(root.xpath("//m:oMath", namespaces=ns))
        text = "".join(root.xpath("//w:t/text()|//m:t/text()", namespaces=ns))
        if not math_count or any(x in text for x in (r"\frac", r"\mathrm", r"\sum")):
            raise RuntimeError("DOCX equation conversion failed")
        image_count = sum(name.startswith("word/media/") for name in archive.namelist())
        return math_count, image_count


def main() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    embedded = internal_inputs(markdown)
    if set(embedded) != {x.resolve() for x in FIGURES}:
        raise RuntimeError("paper figure links do not match the three expected figures")
    for path in (*SCORE_INPUTS, *FIGURES, FIGURE_MANIFEST):
        if not path.is_file():
            raise FileNotFoundError(f"required score or figure input is missing: {path}")
    if OUTPUT.exists() or OUTPUT.is_symlink():
        raise FileExistsError(f"paper output already exists: {OUTPUT}")
    inputs = [SOURCE, *SCORE_INPUTS, *FIGURES, FIGURE_MANIFEST]
    input_hashes = {str(path.relative_to(ROOT)): sha256(path) for path in inputs}
    with tempfile.TemporaryDirectory(prefix="first-hour-paper-build-", dir=EVIDENCE) as staging:
        build = Path(staging) / "paper"
        build.mkdir()
        docx = build / f"{NAME}.docx"
        pdf = build / f"{NAME}.pdf"
        result = run("pandoc", "--from=markdown+tex_math_single_backslash-implicit_figures", "--to=docx",
                     "--resource-path=" + str(ROOT / "docs"), str(SOURCE), "--output", str(docx))
        if result.stderr.strip():
            raise RuntimeError("Pandoc warning: " + result.stderr)
        format_docx(docx)
        math_count, image_count = check_docx(docx)
        if image_count < len(FIGURES):
            raise RuntimeError(f"DOCX embeds {image_count} images, expected {len(FIGURES)}")
        with tempfile.TemporaryDirectory(prefix="first-hour-paper-lo-") as profile:
            run("libreoffice", "-env:UserInstallation=" + Path(profile).as_uri(),
                "--headless", "--convert-to", "pdf", "--outdir", str(build), str(docx))
        pdf_info = run("pdfinfo", str(pdf)).stdout
        pages = next(int(row.split(":", 1)[1]) for row in pdf_info.splitlines()
                     if row.startswith("Pages:"))
        pdf_text = run("pdftotext", str(pdf), "-").stdout
        if "국문 초록" not in pdf_text or "B0" not in pdf_text:
            raise RuntimeError("PDF text check did not find abstract and factorial cases")
        if re.search(r"\\[A-Za-z]+|\\[()\[\]]", pdf_text):
            raise RuntimeError("PDF contains unrendered TeX commands or delimiters")
        if "∨" in pdf_text or "¿" in pdf_text:
            raise RuntimeError("PDF contains corrupted absolute-value math glyphs")
        if "abs(e_i)" not in re.sub(r"\s+", "", pdf_text):
            raise RuntimeError("PDF text check did not retain the absolute-error definition")
        if "MAE" not in pdf_text or "RMSE" not in pdf_text:
            raise RuntimeError("PDF text check did not retain the continuous-score equations")
        if {str(path.relative_to(ROOT)): sha256(path) for path in inputs} != input_hashes:
            raise RuntimeError("paper, score, or figure input changed during conversion")
        receipt = {
            "source": str(SOURCE.relative_to(ROOT)),
            "inputs_sha256": input_hashes,
            "docx_sha256": sha256(docx),
            "pdf_sha256": sha256(pdf),
            "docx_math_objects": math_count,
            "docx_embedded_images": image_count,
            "pdf_pages": pages,
            "paper_format": "A4",
            "font": "NanumMyeongjo",
            "pdf_equation_text_check": "abs(e_i) definition and MAE/RMSE present; corrupted glyphs absent",
            "pandoc": run("pandoc", "--version").stdout.splitlines()[0],
            "libreoffice": run("libreoffice", "--version").stdout.splitlines()[0],
            "note": "Build checks verify link resolution, input hashes, math conversion and embedded images; visual review is separate.",
        }
        (build / "build_receipt.json").write_text(
            json.dumps(receipt, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        build.rename(OUTPUT)
    print(json.dumps(receipt, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
